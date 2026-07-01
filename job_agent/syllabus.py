from __future__ import annotations

import base64
import hashlib
import json
import os
import re
import shutil
import sqlite3
from datetime import datetime
from pathlib import Path
from typing import Any

from .storage import connect


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


def add_syllabus(db_path: Path, project_root: Path, file_path: Path, title: str | None = None) -> sqlite3.Row:
    file_path = file_path.resolve()
    if not file_path.exists():
        raise FileNotFoundError(str(file_path))
    ext = file_path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported syllabus file type: {ext}. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}")

    raw_text = extract_text(file_path)
    if not raw_text.strip():
        raise ValueError(f"No readable text extracted from {file_path}")

    content_hash = hashlib.sha256(raw_text.encode("utf-8")).hexdigest()
    now = datetime.now().isoformat(timespec="seconds")
    inferred_title = title or infer_title(raw_text, file_path)

    data_dir = Path(os.getenv("JOB_AGENT_DATA_DIR") or project_root / "data").resolve()
    stored_dir = data_dir / "syllabi_files"
    stored_dir.mkdir(parents=True, exist_ok=True)
    stored_path = stored_dir / f"{content_hash[:12]}_{file_path.name}"
    if not stored_path.exists():
        shutil.copy2(file_path, stored_path)

    with connect(db_path) as conn:
        existing_by_source = conn.execute(
            "SELECT id FROM syllabi WHERE source_path = ?",
            (str(file_path),),
        ).fetchone()
        if existing_by_source:
            conn.execute(
                """
                UPDATE syllabi
                SET title = ?, file_name = ?, stored_path = ?, file_type = ?,
                    content_hash = ?, raw_text = ?, updated_at = ?
                WHERE id = ?
                """,
                (
                    inferred_title,
                    file_path.name,
                    str(stored_path),
                    ext.lstrip("."),
                    content_hash,
                    raw_text,
                    now,
                    existing_by_source["id"],
                ),
            )
            conn.commit()
            return conn.execute("SELECT * FROM syllabi WHERE id = ?", (existing_by_source["id"],)).fetchone()

        conn.execute(
            """
            INSERT OR IGNORE INTO syllabi (
                title, file_name, source_path, stored_path, file_type,
                content_hash, raw_text, created_at, updated_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                inferred_title,
                file_path.name,
                str(file_path),
                str(stored_path),
                ext.lstrip("."),
                content_hash,
                raw_text,
                now,
                now,
            ),
        )
        conn.execute(
            """
            UPDATE syllabi
            SET title = ?, source_path = ?, stored_path = ?, raw_text = ?, updated_at = ?
            WHERE content_hash = ?
            """,
            (inferred_title, str(file_path), str(stored_path), raw_text, now, content_hash),
        )
        conn.commit()
        return conn.execute("SELECT * FROM syllabi WHERE content_hash = ?", (content_hash,)).fetchone()


def extract_text(file_path: Path) -> str:
    ext = file_path.suffix.lower()
    if ext in {".txt", ".md"}:
        return read_text_file(file_path)
    if ext == ".pdf":
        return extract_pdf_text(file_path)
    if ext == ".docx":
        return extract_docx_text(file_path)
    raise ValueError(f"Unsupported file type: {ext}")


def read_text_file(file_path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return file_path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return file_path.read_text(errors="ignore")


def extract_pdf_text(file_path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(file_path))
    chunks: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            chunks.append(f"--- 第 {index} 页 ---\n{text}")
    extracted = "\n\n".join(chunks)

    enough_text = len(extracted.strip()) >= max(300, len(reader.pages) * 40)
    if enough_text:
        return extracted

    if os.getenv("OPENAI_API_KEY"):
        try:
            return extract_pdf_text_with_pdf_parse_api(file_path, is_ocr=True, end_pages=len(reader.pages))
        except Exception as exc:
            print(f"PDF parse API failed, falling back to vision OCR: {type(exc).__name__}: {exc}", flush=True)

    if os.getenv("OPENAI_API_KEY"):
        return extract_pdf_text_with_vision(file_path)
    return extracted


def extract_pdf_text_with_pdf_parse_api(file_path: Path, is_ocr: bool = True, end_pages: int = 0) -> str:
    """Parse PDFs with the provider's dedicated async PDF/OCR API."""
    import time

    import requests

    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        raise RuntimeError("OPENAI_API_KEY is required for the PDF parse API.")

    base_url = (os.getenv("PDF_PARSE_BASE_URL") or "https://api.gpt.ge").rstrip("/")
    submit_url = f"{base_url}/task/gi/pdf-parse"
    headers = {"Authorization": f"Bearer {api_key}"}
    data = {
        "end_pages": str(end_pages),
        "is_ocr": "true" if is_ocr else "false",
        "language": os.getenv("PDF_PARSE_LANGUAGE", "ch"),
        "formula_enable": os.getenv("PDF_PARSE_FORMULA_ENABLE", "false"),
        "table_enable": os.getenv("PDF_PARSE_TABLE_ENABLE", "true"),
        "layout_model": os.getenv("PDF_PARSE_LAYOUT_MODEL", "doclayout_yolo"),
    }

    print(f"Submitting PDF parse task: {file_path.name}", flush=True)
    with file_path.open("rb") as f:
        response = requests.post(
            submit_url,
            headers=headers,
            data=data,
            files={"file": (file_path.name, f, "application/pdf")},
            timeout=120,
        )
    if not response.ok:
        raise RuntimeError(f"{response.status_code} {response.text[:1200]}")
    payload = response.json()
    task_id = payload.get("task_id")
    if not task_id:
        raise RuntimeError(f"PDF parse task did not return task_id: {payload}")

    poll_url = f"{base_url}/task/{task_id}"
    max_wait_seconds = int(os.getenv("PDF_PARSE_MAX_WAIT_SECONDS", "180"))
    poll_interval = float(os.getenv("PDF_PARSE_POLL_INTERVAL_SECONDS", "5"))
    deadline = time.monotonic() + max_wait_seconds
    attempt = 0

    while time.monotonic() < deadline:
        attempt += 1
        poll_response = requests.get(poll_url, headers=headers, timeout=60)
        poll_response.raise_for_status()
        result = poll_response.json()
        status = result.get("status")
        data_obj = result.get("data") if isinstance(result.get("data"), dict) else {}
        state = data_obj.get("state")
        progress = data_obj.get("progress")
        print(f"PDF parse poll {attempt}: status={status}, state={state}, progress={progress}", flush=True)

        if status == "success" or state == 1:
            return parse_pdf_task_output(result)
        if isinstance(state, int) and state < 0:
            raise RuntimeError(f"PDF parse task failed: {result}")
        time.sleep(poll_interval)

    raise TimeoutError(f"PDF parse task timed out after {max_wait_seconds}s: {task_id}")


def parse_pdf_task_output(payload: dict[str, Any]) -> str:
    output = payload.get("output")
    if isinstance(output, dict):
        segments = output.get("segments")
        if isinstance(segments, list):
            chunks: list[str] = []
            for segment in segments:
                if not isinstance(segment, dict):
                    continue
                index = segment.get("index")
                content = str(segment.get("content") or "").strip()
                if content:
                    chunks.append(f"--- 第 {index} 页 ---\n{content}")
            if chunks:
                return "\n\n".join(chunks)

    text = payload.get("text")
    if isinstance(text, str) and text.strip():
        return text.strip()
    raise RuntimeError(f"PDF parse task returned no text: {payload}")


def extract_pdf_text_with_vision(file_path: Path) -> str:
    """Fallback OCR for scanned PDFs using the configured OpenAI-compatible vision model."""
    import fitz
    from openai import OpenAI

    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        raise RuntimeError("This PDF appears to be scanned. OPENAI_API_KEY is required for OCR fallback.")

    client = OpenAI(api_key=api_key, base_url=os.getenv("OPENAI_BASE_URL") or None)
    model = os.getenv("OCR_MODEL") or os.getenv("OPENAI_MODEL", "gpt-4o-mini")

    doc = fitz.open(str(file_path))
    chunks: list[str] = []
    try:
        for page_index in range(doc.page_count):
            page = doc.load_page(page_index)
            pix = page.get_pixmap(matrix=fitz.Matrix(1.8, 1.8), alpha=False)
            image_bytes = pix.tobytes("png")
            data_url = "data:image/png;base64," + base64.b64encode(image_bytes).decode("ascii")
            print(f"OCR PDF page {page_index + 1}/{doc.page_count}: {file_path.name}", flush=True)
            response = client.chat.completions.create(
                model=model,
                temperature=0,
                messages=[
                    {
                        "role": "system",
                        "content": (
                            "你是教学大纲 OCR 助手。请逐字识别图片中的中文和英文文本，"
                            "保留标题、段落、表格中的关键信息。不要总结，不要添加图片中没有的内容。"
                        ),
                    },
                    {
                        "role": "user",
                        "content": [
                            {"type": "text", "text": "请识别这一页教学大纲中的所有可读文字。"},
                            {"type": "image_url", "image_url": {"url": data_url}},
                        ],
                    },
                ],
            )
            text = (response.choices[0].message.content or "").strip()
            chunks.append(f"--- 第 {page_index + 1} 页 ---\n{text}")
    finally:
        doc.close()
    return "\n\n".join(chunks)


def extract_docx_text(file_path: Path) -> str:
    from docx import Document

    doc = Document(str(file_path))
    chunks: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            chunks.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                chunks.append(" | ".join(cells))
    return "\n".join(chunks)


def infer_title(raw_text: str, file_path: Path) -> str:
    course_name = infer_course_name(raw_text)
    if course_name:
        return course_name
    for line in raw_text.splitlines():
        line = re.sub(r"^#+\s*", "", line.strip())
        if not line or re.match(r"^-+\s*第\s*\d+\s*页\s*-+$", line):
            continue
        if len(line) <= 80:
            return line
    return file_path.stem


def infer_course_name(raw_text: str) -> str | None:
    patterns = [
        r"课程名称\s*(?:\(Course\)|（Course）)?\s*[:：]\s*([^\n\r|]+)",
        r"课程\s*(?:\(Course\)|（Course）)?\s*[:：]\s*([^\n\r|]+)",
        r"Course\s*[:：]\s*([^\n\r|]+)",
    ]
    for pattern in patterns:
        match = re.search(pattern, raw_text, flags=re.IGNORECASE)
        if not match:
            continue
        name = re.sub(r"\s+", " ", match.group(1)).strip(" ：:;；")
        name = re.sub(r"\s{2,}.*$", "", name).strip()
        if 1 <= len(name) <= 60:
            return name
    return None


def list_syllabi(db_path: Path) -> list[sqlite3.Row]:
    with connect(db_path) as conn:
        return conn.execute(
            """
            SELECT id, title, file_name, file_type, length(raw_text) AS text_length, updated_at
            FROM syllabi
            ORDER BY id DESC
            """
        ).fetchall()


def get_syllabus(db_path: Path, syllabus_id: int) -> sqlite3.Row | None:
    with connect(db_path) as conn:
        return conn.execute("SELECT * FROM syllabi WHERE id = ?", (syllabus_id,)).fetchone()


def get_syllabi(db_path: Path, syllabus_ids: list[int]) -> list[sqlite3.Row]:
    if not syllabus_ids:
        return []
    placeholders = ",".join("?" for _ in syllabus_ids)
    with connect(db_path) as conn:
        return conn.execute(f"SELECT * FROM syllabi WHERE id IN ({placeholders})", syllabus_ids).fetchall()


def analyze_syllabus(db_path: Path, syllabus_id: int) -> sqlite3.Row:
    from openai import OpenAI

    syllabus = get_syllabus(db_path, syllabus_id)
    if not syllabus:
        raise ValueError(f"Syllabus not found: #{syllabus_id}")
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        raise RuntimeError("OPENAI_API_KEY is not configured.")

    client = OpenAI(api_key=api_key, base_url=os.getenv("OPENAI_BASE_URL") or None)
    model = os.getenv("OPENAI_MODEL", "gpt-4o-mini")
    response = client.chat.completions.create(
        model=model,
        temperature=0.2,
        messages=[
            {
                "role": "system",
                "content": (
                    "你是课程教学大纲分析助手。请只依据大纲原文分析课程教给学生的知识、能力、"
                    "技术工具方法和可能相关岗位方向。不要编造大纲没有体现的内容。只返回合法 JSON。"
                ),
            },
            {"role": "user", "content": build_profile_prompt(syllabus)},
        ],
    )
    profile = parse_json_response(response.choices[0].message.content or "")
    return save_syllabus_profile(db_path, syllabus_id, profile)


def build_profile_prompt(syllabus: sqlite3.Row) -> str:
    text = str(syllabus["raw_text"])[:24000]
    return (
        f"课程大纲标题：{syllabus['title']}\n"
        f"文件名：{syllabus['file_name']}\n\n"
        "请输出 JSON，格式如下：\n"
        "{\n"
        '  "course_name": "课程名称",\n'
        '  "summary": "课程简要说明",\n'
        '  "knowledge_points": ["学生掌握的知识点"],\n'
        '  "abilities": ["学生具备的能力"],\n'
        '  "technologies_tools_methods": ["学生接触的技术、工具、方法"],\n'
        '  "job_directions": ["这些能力可能对应的岗位方向"],\n'
        '  "evidence": ["来自大纲原文的关键依据"]\n'
        "}\n\n"
        "大纲原文：\n"
        f"{text}"
    )


def save_syllabus_profile(db_path: Path, syllabus_id: int, profile: dict[str, Any]) -> sqlite3.Row:
    now = datetime.now().isoformat(timespec="seconds")
    summary = str(profile.get("summary") or "")
    course_name = str(profile.get("course_name") or "").strip()
    with connect(db_path) as conn:
        conn.execute(
            """
            INSERT INTO syllabus_profiles (
                syllabus_id, summary, knowledge_points_json, abilities_json,
                technologies_tools_methods_json, job_directions_json,
                profile_json, created_at, updated_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(syllabus_id) DO UPDATE SET
                summary = excluded.summary,
                knowledge_points_json = excluded.knowledge_points_json,
                abilities_json = excluded.abilities_json,
                technologies_tools_methods_json = excluded.technologies_tools_methods_json,
                job_directions_json = excluded.job_directions_json,
                profile_json = excluded.profile_json,
                updated_at = excluded.updated_at
            """,
            (
                syllabus_id,
                summary,
                json.dumps(profile.get("knowledge_points", []), ensure_ascii=False),
                json.dumps(profile.get("abilities", []), ensure_ascii=False),
                json.dumps(profile.get("technologies_tools_methods", []), ensure_ascii=False),
                json.dumps(profile.get("job_directions", []), ensure_ascii=False),
                json.dumps(profile, ensure_ascii=False),
                now,
                now,
            ),
        )
        if course_name:
            conn.execute(
                "UPDATE syllabi SET title = ?, updated_at = ? WHERE id = ?",
                (course_name, now, syllabus_id),
            )
        conn.commit()
        return conn.execute("SELECT * FROM syllabus_profiles WHERE syllabus_id = ?", (syllabus_id,)).fetchone()


def get_syllabus_profile(db_path: Path, syllabus_id: int) -> sqlite3.Row | None:
    with connect(db_path) as conn:
        return conn.execute("SELECT * FROM syllabus_profiles WHERE syllabus_id = ?", (syllabus_id,)).fetchone()


def parse_json_response(raw: str) -> dict[str, Any]:
    cleaned = raw.strip()
    if cleaned.startswith("```json"):
        cleaned = cleaned.removeprefix("```json").strip()
    if cleaned.startswith("```"):
        cleaned = cleaned.removeprefix("```").strip()
    if cleaned.endswith("```"):
        cleaned = cleaned[: -3].strip()
    return json.loads(cleaned)
